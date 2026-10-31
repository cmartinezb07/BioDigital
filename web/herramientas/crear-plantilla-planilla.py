#!/usr/bin/env python3
"""Crea la planilla .xlsx del dueño a partir de los 4 CSV de un negocio.

Uso:
  python web/herramientas/crear-plantilla-planilla.py --prefijo barberia --nombre "Barbería El Filo" \
      --color 8e0f0f --oscuro 141414 --whatsapp 56946941640

Lee   web/propuesta/ejemplos/reservas-simple/planilla-ejemplo/<prefijo>-{servicios,horario,barberos,aviso}.csv
Crea  docs/plantillas/planilla-<prefijo>.xlsx

Regla dura: la fila 1 de cada pestaña de datos es el encabezado EXACTO del CSV (el sitio lee por nombre de columna).
Sin banners, sin columnas extra, sin celdas combinadas en los datos. Las ayudas van en notas del encabezado y en "Léeme".
"""
import argparse, csv, datetime as dt, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parents[2]
CARPETA_CSV = RAIZ / "web/propuesta/ejemplos/reservas-simple/planilla-ejemplo"
DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
HORAS = [f"{h:02d}:{m:02d}" for h in range(6, 24) for m in (0, 15, 30, 45)]
FILAS = 30  # filas con formato listas para agregar más (incluye el encabezado)
MAXF = 200  # última fila de datos desbloqueada (validaciones y formato condicional llegan hasta aquí)
FUENTE = "Calibri"

# Qué hace cada pestaña (para "Léeme"); si el negocio trae otra pestaña, se muestra con su nombre.
DESCRIPCION = {
    "servicios": ("Servicios", "Tus servicios, precios y duración. Cambia un precio, agrega un servicio o apágalo con Activo = No."),
    "horario": ("Horario", "Los días y horas en que atiendes. Si cierras al almuerzo, escribe dos filas del mismo día."),
    "barberos": ("Barberos", "Quiénes atienden. Activo = No lo oculta (por ejemplo, si está de vacaciones)."),
    "aviso": ("Aviso", "Un mensaje destacado en tu página (promoción, feriado). Se ve solo entre las fechas Desde y Hasta."),
}
AYUDA = {  # nota en el encabezado de cada columna
    "servicio": "Nombre del servicio, como lo verá tu cliente. Ejemplo: Fade.",
    "categoría": "Elige de la lista: agrupa el servicio en su pestaña de la página.",
    "precio": "Solo números, sin $ ni puntos. Ejemplo: 15000.",
    "duración (min)": "Minutos que dura. Solo números. Define cuántas horas se ofrecen. Ejemplo: 45.",
    "descripción": "Opcional. Una frase corta (se puede dejar vacía).",
    "activo": "Elige Sí para mostrarlo o No para ocultarlo sin borrar la fila.",
    "día": "Elige el día de la lista.",
    "abre": "Elige la hora de la lista (de 15 en 15 minutos). Ejemplo: 10:00.",
    "cierra": "Elige la hora de la lista. Debe ser después de la hora de apertura.",
    "barbero": "Nombre como quieres que aparezca. Ejemplo: Cuervo.",
    "título": "Título corto del aviso. Ejemplo: Semana de descuento.",
    "texto": "El mensaje completo del aviso (máximo 2 o 3 líneas).",
    "desde": "Primer día en que se muestra. Formato DD-MM-AAAA, ejemplo: 28-10-2026.",
    "hasta": "Último día en que se muestra. Formato DD-MM-AAAA, ejemplo: 30-11-2026.",
}
ANCHO = {"servicio": 28, "categoría": 16, "precio": 12, "duración (min)": 16, "descripción": 48, "activo": 11, "día": 16,
         "abre": 12, "cierra": 12, "barbero": 26, "título": 28, "texto": 60, "desde": 15, "hasta": 15}


def rgb(h):
    return h.lstrip("#").upper()


def leer_csv(ruta):
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        filas = [r for r in csv.reader(f)]
    return filas[0], [r for r in filas[1:] if any(c.strip() for c in r)]


def proteger(ws):
    # En openpyxl, True = prohibido. Se permite: seleccionar, ordenar, autofiltro, insertar/eliminar filas, ancho/alto.
    p = ws.protection
    p.sheet = True
    p.formatCells = True        # no cambiar formato de celdas
    p.objects = True            # no tocar objetos
    p.scenarios = True
    p.selectLockedCells = False
    p.selectUnlockedCells = False
    p.sort = False
    p.autoFilter = False
    p.insertRows = False
    p.deleteRows = False
    p.formatColumns = False
    p.formatRows = False
    p.insertColumns = True
    p.deleteColumns = True
    p.insertHyperlinks = True
    p.pivotTables = True


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--prefijo", required=True)
    ap.add_argument("--nombre", required=True)
    ap.add_argument("--color", default="8e0f0f", help="color de marca (hex, para encabezados)")
    ap.add_argument("--oscuro", default="141414", help="color oscuro (banner de Léeme)")
    ap.add_argument("--whatsapp", required=True, help="número con código de país, ej. 56946941640")
    ap.add_argument("--carpeta-csv", default=str(CARPETA_CSV))
    ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    marca, oscuro = rgb(a.color), rgb(a.oscuro)

    archivos = sorted(Path(a.carpeta_csv).glob(f"{a.prefijo}-*.csv"))
    if not archivos:
        sys.exit(f"No hay CSV '{a.prefijo}-*.csv' en {a.carpeta_csv}")
    orden = ["servicios", "horario", "barberos", "aviso"]
    archivos.sort(key=lambda p: orden.index(p.stem.split("-", 1)[1]) if p.stem.split("-", 1)[1] in orden else 99)

    wb = Workbook()
    lee = wb.active
    lee.title = "Léeme"
    listas = wb.create_sheet("Listas")  # horas de las listas desplegables (oculta)
    for i, h in enumerate(HORAS, 1):
        listas.cell(i, 1, h).number_format = "@"
    listas.sheet_state = "hidden"
    rango_horas = f"Listas!$A$1:$A${len(HORAS)}"

    fino = Side(style="thin", color="D0D0D0")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    colores_tab = [marca, "2F6F4F", "B07A1F", "3A5E8C", "6B4B8A"]
    hojas = []

    for n, ruta in enumerate(archivos):
        clave = ruta.stem.split("-", 1)[1]
        titulo = DESCRIPCION.get(clave, (clave.capitalize(), ""))[0]
        enc, filas = leer_csv(ruta)
        ws = wb.create_sheet(titulo, index=1 + n)
        hojas.append((clave, titulo, enc))
        ws.sheet_properties.tabColor = colores_tab[(n + 1) % len(colores_tab)]
        ws.sheet_view.zoomScale = 110
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 30
        k = [c.strip().lower() for c in enc]
        for j, nombre in enumerate(enc, 1):
            c = ws.cell(1, j, nombre)  # nombre EXACTO del CSV
            c.fill = PatternFill("solid", fgColor=marca)
            c.font = Font(name=FUENTE, bold=True, color="FFFFFF", size=11)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = borde
            ayuda = AYUDA.get(k[j - 1], "")
            if ayuda:
                cm = Comment(ayuda, "Bío Digital")
                cm.width, cm.height = 280, 90
                c.comment = cm
            ws.column_dimensions[L(j)].width = ANCHO.get(k[j - 1], max(14, min(40, len(nombre) + 6)))

        es_fecha = lambda t: t in ("desde", "hasta") and clave == "aviso"
        es_hora = lambda t: t in ("abre", "cierra")
        largo = any(t in ("texto", "descripción") for t in k)
        for r in range(2, FILAS + 1):
            ejemplo = r - 2 < len(filas)
            if ejemplo:
                fondo = "EDEDED" if r % 2 == 0 else "F6F6F6"   # gris = fila de ejemplo
            else:
                fondo = "FFFFFF" if r % 2 == 0 else "FAFAFA"   # blanco = libre para ti
            for j in range(1, len(enc) + 1):
                t = k[j - 1]
                c = ws.cell(r, j)
                c.fill = PatternFill("solid", fgColor=fondo)
                c.border = borde
                c.font = Font(name=FUENTE, size=11, color="222222")
                h, wrap = "left", False
                if ejemplo:
                    v = filas[r - 2][j - 1].strip() if j - 1 < len(filas[r - 2]) else ""
                    if es_fecha(t) and v:
                        d, m, y = (int(x) for x in v.replace("/", "-").split("-"))
                        c.value = dt.datetime(y, m, d)
                    elif t in ("precio", "duración (min)") and v.isdigit():
                        c.value = int(v)
                    elif v != "":
                        c.value = v
                if t in ("precio", "duración (min)"):
                    c.number_format = "0"
                    h = "right"
                elif es_fecha(t):
                    c.number_format = "DD-MM-YYYY"
                    h = "center"
                elif es_hora(t):
                    c.number_format = "@"
                    h = "center"
                elif t == "activo":
                    h = "center"
                elif t in ("texto", "descripción"):
                    wrap = True
                c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap, indent=1 if h == "left" else 0)
            ws.row_dimensions[r].height = 36 if largo else 24

        # Datos desbloqueados (filas 2 a MAXF); encabezado y todo lo demás queda bloqueado. Hoja protegida SIN contraseña.
        for r in range(2, MAXF + 1):
            for j in range(1, len(enc) + 1):
                ws.cell(r, j).protection = Protection(locked=False)
        proteger(ws)

        ult = len(enc)

        def validacion(j, **kw):
            dv = DataValidation(allow_blank=True, showErrorMessage=True, showInputMessage=True, **kw)
            dv.add(f"{L(j)}2:{L(j)}{MAXF}")
            ws.add_data_validation(dv)
            return dv

        for i, t in enumerate(k):
            j = i + 1
            if t == "activo":
                d = validacion(j, type="list", formula1='"Sí,No"', errorStyle="warning")
                d.errorTitle, d.error = "Activo", "Escribe Sí o No (elige de la lista)."
                d.promptTitle, d.prompt = "Activo", "Sí = se muestra · No = se oculta."
                col = L(j)
                rng = f"{col}2:{col}{MAXF}"
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({col}2="No",{col}2="NO",{col}2="no")'], font=Font(bold=True, color="9C0006"), fill=PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")))
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({col}2="Sí",{col}2="SÍ",{col}2="Si",{col}2="SI",{col}2="sí",{col}2="si")'], font=Font(bold=True, color="006100"), fill=PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")))
                if ult > 1:
                    ws.conditional_formatting.add(f"A2:{L(ult)}{MAXF}", FormulaRule(formula=[f'OR(${col}2="No",${col}2="NO",${col}2="no")'], font=Font(color="9A9A9A", italic=True), fill=PatternFill("solid", bgColor="D9D9D9", fgColor="D9D9D9")))
            elif t == "categoría":
                vistas = sorted({r[i].strip() for r in filas if len(r) > i and r[i].strip()})
                base = ["Cortes", "Barba", "Otros"]
                lista = base + [x for x in vistas if x not in base]
                d = validacion(j, type="list", formula1='"' + ",".join(lista) + '"', errorStyle="warning")
                d.errorTitle, d.error = "Categoría", "Elige una categoría de la lista para que el servicio quede bien agrupado."
                d.promptTitle, d.prompt = "Categoría", "Elige de la lista."
            elif t == "día":
                d = validacion(j, type="list", formula1='"' + ",".join(DIAS) + '"', errorStyle="warning")
                d.errorTitle, d.error = "Día", "Elige un día de la lista (Lunes a Domingo)."
                d.promptTitle, d.prompt = "Día", "Elige de la lista."
            elif es_hora(t):
                d = validacion(j, type="list", formula1=f"={rango_horas}", errorStyle="warning")
                d.errorTitle, d.error = "Hora", "Elige la hora de la lista, con este formato: 09:30."
                d.promptTitle, d.prompt = "Hora", "Elige de la lista (cada 15 min). Formato 09:30."
            elif t == "precio":
                d = validacion(j, type="whole", operator="between", formula1="0", formula2="500000", errorStyle="stop")
                d.errorTitle, d.error = "Precio fuera de rango", "Escribe el precio en pesos, sin puntos ni símbolo. Ejemplo: 15000. El máximo es 500.000."
                d.promptTitle, d.prompt = "Precio en pesos", "Solo números, sin puntos ni $. Ejemplo: 15000 (máximo 500.000)."
                col = L(j)
                rng = f"{col}2:{col}{MAXF}"
                rojo = PatternFill("solid", bgColor="FF0000", fgColor="FF0000")
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({col}2<>"",OR(NOT(ISNUMBER({col}2)),{col}2<0,{col}2>500000))'], font=Font(bold=True, color="FFFFFF"), fill=rojo, stopIfTrue=True))
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({col}2),{col}2>100000)'], font=Font(bold=True, color="7F6000"), fill=PatternFill("solid", bgColor="FFEB9C", fgColor="FFEB9C")))
            elif t == "duración (min)":
                d = validacion(j, type="whole", operator="between", formula1="5", formula2="720", errorStyle="stop")
                d.errorTitle, d.error = "Duración fuera de rango", "Escribe los minutos como número entero, de 5 a 720. Ejemplo: 45."
                d.promptTitle, d.prompt = "Duración en minutos", "Minutos, solo números. Ejemplo: 45 (entre 5 y 720)."
            elif t in ("servicio", "barbero", "título", "texto", "descripción"):
                mx, est, ej, nom = {"servicio": (60, "warning", "Fade", "Nombre"), "barbero": (40, "warning", "Cuervo", "Nombre"),
                                    "título": (60, "stop", "Semana de descuento", "Título"), "texto": (200, "warning", "15% en Corte + barba de martes a jueves.", "Texto"),
                                    "descripción": (150, "warning", "Degradado a máquina y terminación a tijera.", "Descripción")}[t]
                d = validacion(j, type="textLength", operator="lessThanOrEqual", formula1=str(mx), errorStyle=est)
                d.errorTitle, d.error = f"{nom} muy largo", f"Máximo {mx} caracteres. Acórtalo para que se vea bien en tu página."
                d.promptTitle, d.prompt = nom, f"Ejemplo: {ej} (máximo {mx} caracteres)."
            elif es_fecha(t):
                d = validacion(j, type="date", operator="greaterThan", formula1="43831", errorStyle="warning")  # 01-01-2020
                d.errorTitle, d.error = "Fecha", "Escribe la fecha como DD-MM-AAAA. Ejemplo: 28-10-2026."
                d.promptTitle, d.prompt = "Fecha", "Formato DD-MM-AAAA. Ejemplo: 28-10-2026."

        # ---- Avisos por formato condicional (rojo = corregir, amarillo = revisar). Sin columnas extra. ----
        idx = {t: L(i + 1) for i, t in enumerate(k)}
        ROJO = dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", bgColor="FF0000", fgColor="FF0000"))
        AMAR = dict(font=Font(bold=True, color="7F6000"), fill=PatternFill("solid", bgColor="FFEB9C", fgColor="FFEB9C"))

        def regla(rango, formula, estilo):
            ws.conditional_formatting.add(rango, FormulaRule(formula=[formula], **estilo))

        def R(c):
            return f"{idx[c]}2:{idx[c]}{MAXF}"
        if clave == "horario" and {"día", "abre", "cierra"} <= set(idx):
            D, A_, C_ = idx["día"], idx["abre"], idx["cierra"]
            # Se comparan VALORES de hora (no texto): "9:00", "09:00" y "10:00:00" valen igual; una hora no interpretable da -1 y se marca roja.
            hv = lambda c: f'IF(ISNUMBER(${c}2),${c}2,IFERROR(TIMEVALUE(${c}2),-1))'
            regla(f"{A_}2:{C_}{MAXF}", f'AND(OR(${A_}2<>"",${C_}2<>""),OR(${A_}2="",${C_}2="",{hv(A_)}<0,{hv(C_)}<0,{hv(C_)}<={hv(A_)}))', ROJO)
            regla(f"{D}2:{C_}{MAXF}", f'AND(${D}2<>"",COUNTIFS(${D}$2:${D}${MAXF},${D}2,${A_}$2:${A_}${MAXF},${A_}2)>1)', AMAR)
        if clave == "aviso" and {"título", "texto", "desde", "hasta"} <= set(idx):
            T, X, De, Ha = idx["título"], idx["texto"], idx["desde"], idx["hasta"]
            regla(f"{De}2:{Ha}{MAXF}", f'AND(ISNUMBER(${De}2),ISNUMBER(${Ha}2),${Ha}2<${De}2)', ROJO)
            regla(R("título"), f'LEN({T}2)>60', ROJO)
            regla(R("texto"), f'LEN({X}2)>200', AMAR)
        if clave == "servicios" and {"servicio", "precio"} <= set(idx):
            S, P = idx["servicio"], idx["precio"]
            falta = f'OR(AND(${S}2<>"",${P}2=""),AND(${S}2="",${P}2<>""))'
            regla(R("servicio"), f'OR(AND({S}2<>"",COUNTIF(${S}$2:${S}${MAXF},{S}2)>1),{falta})', ROJO)
            regla(R("precio"), falta, ROJO)
            if "descripción" in idx:
                regla(R("descripción"), f'LEN({idx["descripción"]}2)>150', AMAR)
        if clave == "barberos" and "barbero" in idx:
            B = idx["barbero"]
            regla(R("barbero"), f'AND({B}2<>"",COUNTIF(${B}$2:${B}${MAXF},{B}2)>1)', ROJO)

    # ---------- Léeme ----------
    lee.sheet_properties.tabColor = oscuro
    lee.sheet_view.zoomScale = 110
    lee.sheet_view.showGridLines = False
    for col, w in (("A", 3), ("B", 8), ("C", 34), ("D", 78)):
        lee.column_dimensions[col].width = w

    def f(**kw):
        return Font(name=FUENTE, **kw)

    def banner(fila, texto, fondo, size=12, alto=26):
        lee.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
        c = lee.cell(fila, 2, texto)
        c.font = f(bold=True, color="FFFFFF", size=size)
        c.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
        for col in (2, 3, 4):
            lee.cell(fila, col).fill = PatternFill("solid", fgColor=fondo)
        lee.row_dimensions[fila].height = alto

    def texto(fila, col, t, **kw):
        c = lee.cell(fila, col, t)
        c.font = f(**kw)
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1 if col == 3 else 0)
        return c

    banner(2, a.nombre, oscuro, size=24, alto=52)
    banner(3, "Tu planilla: cambia aquí tus precios, horarios y avisos y tu página se actualiza sola.", marca, size=12, alto=30)
    r = 5
    banner(r, "Cómo se usa, en 3 pasos", marca, size=13, alto=28)
    r += 1
    pasos = [("1", "Abre la pestaña que quieras cambiar", "Están abajo, junto a esta: " + ", ".join(h[1] for h in hojas) + "."),
             ("2", "Cambia solo las celdas blancas o grises", "Escribe encima del dato. Donde veas una flechita, elige de la lista. Para agregar algo nuevo, usa la primera fila vacía."),
             ("3", "Mira el cambio en tu página", "Recarga tu página (puede tardar unos minutos en aparecer). No necesitas guardar: la planilla se guarda sola.")]
    for num, t, d in pasos:
        n = lee.cell(r, 2, num)
        n.fill = PatternFill("solid", fgColor=marca)
        n.font = f(bold=True, color="FFFFFF", size=18)
        n.alignment = Alignment(horizontal="center", vertical="center")
        texto(r, 3, t, bold=True, size=12)
        texto(r, 4, d, size=11, color="333333")
        for col in (2, 3, 4):
            lee.cell(r, col).border = Border(bottom=fino)
        lee.row_dimensions[r].height = 46
        r += 1
    r += 1
    banner(r, "Qué significan los colores", marca, size=13, alto=28)
    r += 1
    leyenda = [(marca, "FFFFFF", "Encabezado de color", "No lo toques: es el nombre de la columna y la página lo necesita tal cual."),
               ("FFFFFF", "222222", "Celda blanca", "Es tuya: escribe aquí lo que quieras."),
               ("EDEDED", "222222", "Fila gris", "Es un ejemplo. Puedes cambiarlo por tus datos o borrarlo."),
               ("C6EFCE", "006100", "Activo = Sí (verde)", "Ese dato se ve en tu página."),
               ("FFC7CE", "9C0006", "Activo = No (rojo)", "Ese dato está oculto. La fila se pone gris para que lo notes."),
               ("FFEB9C", "7F6000", "Precio amarillo", "Mayor a $100.000: revisa que no sobren ceros."),
               ("FF0000", "FFFFFF", "Precio rojo fuerte", "No es un número, es negativo o pasa de $500.000: corrígelo, la página no lo mostraría bien.")]
    for fondo, tc, t, d in leyenda:
        c = texto(r, 3, t, bold=True, color=tc, size=11)
        c.fill = PatternFill("solid", fgColor=fondo)
        c.border = borde
        texto(r, 4, d, size=11, color="333333")
        lee.row_dimensions[r].height = 30
        r += 1
    r += 1
    banner(r, "Avisos de la planilla: rojo = corregir, amarillo = revisar", marca, size=13, alto=28)
    r += 1
    avisos = [("FF0000", "FFFFFF", "Servicios", "Rojo: servicio repetido, o precio sin servicio (o servicio sin precio), o precio fuera de 0 a 500.000. Amarillo: precio sobre 100.000 o descripción de más de 150 caracteres."),
              ("FF0000", "FFFFFF", "Horario", "Rojo: Cierra no es posterior a Abre, o falta una de las dos horas. Amarillo: el mismo día con la misma hora de apertura repetido."),
              ("FF0000", "FFFFFF", "Barberos", "Rojo: nombre repetido."),
              ("FF0000", "FFFFFF", "Aviso", "Rojo: Hasta es anterior a Desde, o el título pasa de 60 caracteres. Amarillo: el texto pasa de 200 caracteres.")]
    for fondo, tc, t, d in avisos:
        c = texto(r, 3, t, bold=True, color=tc, size=11)
        c.fill = PatternFill("solid", fgColor=fondo)
        c.border = borde
        texto(r, 4, d, size=11, color="333333")
        lee.row_dimensions[r].height = 46
        r += 1
    r += 1
    banner(r, "Qué SÍ puedes cambiar", "2F6F4F", size=13, alto=28)
    r += 1
    for clave, titulo, _ in hojas:
        desc = DESCRIPCION.get(clave, (titulo, ""))[1] or "Los datos de esta pestaña."
        texto(r, 3, titulo, bold=True, size=11)
        texto(r, 4, desc, size=11, color="333333")
        for col in (3, 4):
            lee.cell(r, col).border = Border(bottom=fino)
        lee.row_dimensions[r].height = 32
        r += 1
    r += 1
    banner(r, "Qué NO se cambia", "9C0006", size=13, alto=28)
    r += 1
    for t in ["El diseño de la página y su estructura: eso lo hacemos nosotros.",
              "Los nombres de las columnas (primera fila de cada pestaña): si los cambias, la página deja de entender la planilla.",
              "No insertes columnas, no combines celdas, no pongas títulos encima de la tabla ni cambies el nombre de las pestañas."]:
        lee.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        texto(r, 3, "•  " + t, size=11, color="333333")
        lee.row_dimensions[r].height = 32
        r += 1
    r += 1
    banner(r, "¿Algo no funciona?", oscuro, size=13, alto=28)
    r += 1
    lee.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    c = texto(r, 3, f"Escríbenos por WhatsApp: +{a.whatsapp}  (envía una captura de la pestaña y dinos qué quisiste cambiar)", bold=True, size=12, color=marca, underline="single")
    c.hyperlink = f"https://wa.me/{a.whatsapp}"
    lee.row_dimensions[r].height = 36

    lee.merge_cells(start_row=r + 2, start_column=2, end_row=r + 2, end_column=4)
    c = lee.cell(r + 2, 2, "Solo las celdas blancas o grises se editan. Todo lo demás está bloqueado a propósito para no romper tu página.")
    c.font = f(italic=True, size=11, color="555555")
    c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    lee.row_dimensions[r + 2].height = 30
    proteger(lee)
    proteger(listas)
    wb.active = 0
    salida = Path(a.salida) if a.salida else RAIZ / f"docs/plantillas/planilla-{a.prefijo}.xlsx"
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.properties.creator = "Bío Digital"
    wb.properties.lastModifiedBy = "Bío Digital"
    wb.properties.title = f"Planilla del dueño - {a.nombre}"
    wb.save(salida)
    print("Creado", salida)


if __name__ == "__main__":
    main()
